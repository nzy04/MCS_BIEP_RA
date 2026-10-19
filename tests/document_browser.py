import pathlib,sys,json,base64,zipfile,os
WORK=pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0,str(WORK/'test-results/python-deps'))
from playwright.sync_api import sync_playwright,expect
from docx import Document
from docx.shared import Inches
from PIL import Image,ImageDraw
from reportlab.pdfgen import canvas
OUT=WORK/'test-results/document-upgrade';OUT.mkdir(parents=True,exist_ok=True)
BASE=os.environ.get('TEST_BASE_URL','http://127.0.0.1:5190/')
figure=Image.new('RGB',(720,220),'white');d=ImageDraw.Draw(figure)
for x,t in [(20,'Observation'),(270,'Prediction'),(520,'Evaluation')]:d.rectangle((x,60,x+175,155),outline='#173e66',width=3);d.text((x+25,102),t,fill='#173e66')
figure.save(OUT/'figure.png');Image.new('RGB',(120,80),'#43729a').save(OUT/'second.png')
doc=Document();doc.add_heading('图文改造验收笔记',0);doc.add_heading('研究问题与已有工作不足',1);doc.add_paragraph('保留原有章节结构，不按固定栏目重新排列。')
doc.add_heading('框架图与工作流',1);doc.add_paragraph('Observation -> Prediction -> Evaluation，以下为功能测试示意。');doc.add_picture(str(OUT/'figure.png'),width=Inches(5));doc.add_paragraph('图1 原文流程示意，功能测试。')
doc.add_heading('实验分析',1);table=doc.add_table(rows=3,cols=3)
for row,values in zip(table.rows,[['方法','误差','说明'],['基线','0.42','测试'],['预测','0.31','测试']]):
    for c,v in zip(row.cells,values):c.text=v
doc.add_paragraph('合并单元格表格：');table=doc.add_table(rows=2,cols=3);table.cell(0,0).text='指标';table.cell(0,1).merge(table.cell(0,2)).text='合并表头';table.cell(1,0).text='误差';table.cell(1,1).text='0.42';table.cell(1,2).text='0.31'
doc.add_heading('个人判断与论文关联',1);doc.add_paragraph('需要核对长期预测中的误差传播，这段备注不能被自动改写。');doc.save(OUT/'note.docx')
md='# Markdown 图文配套验收\n\n## 自定义第一节\n原始顺序 FIRST。\n\n![框架图](images/figure.png)\n\n图1 图片应保留在第一节。\n\n## 自定义第二节\n原始顺序 SECOND。\n\n| 方法 | 误差 |\n| --- | --- |\n| 基线 | 0.42 |\n| 改进 | 0.31 |\n\n<script>window.__BAD__=true</script>\n'
(OUT/'note.md').write_text(md,encoding='utf-8')
with zipfile.ZipFile(OUT/'note.zip','w') as z:z.writestr('note.md',md);z.write(OUT/'figure.png','images/figure.png')
pdf=canvas.Canvas(str(OUT/'paper.pdf'));pdf.drawString(50,740,'PDF page image import verification. Original paper stays separate from notes.');pdf.showPage();pdf.drawString(50,740,'Second PDF page with a test framework diagram.');pdf.rect(50,400,300,200);pdf.showPage();pdf.save()
results=[]
def demo(route):
    path=route.request.url.split('/api/',1)[1]
    if path=='config':route.fulfill(json={'title':'MCS_BIEP','configured':False,'authConfigured':False,'repository':'','branch':'main'})
    elif path=='session':route.fulfill(status=401,json={'error':'Not signed in'})
    else:route.fulfill(status=404,json={'error':'not configured'})
with sync_playwright() as p:
    browser=p.chromium.launch(channel=os.environ.get('TEST_BROWSER_CHANNEL', 'msedge') or None,headless=True)
    context=browser.new_context(viewport={'width':1440,'height':1100});context.route('**/api/**',demo)
    page=context.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    def start(files,paper=None):
        page.goto(BASE+'#import');page.reload();page.wait_for_load_state('networkidle')
        page.get_by_test_id('note-file').set_input_files([str(f) for f in files]);page.get_by_text('已解析，可继续整理',exact=True).wait_for()
        if paper:page.get_by_test_id('paper-file').set_input_files(str(paper));expect(page.get_by_text('论文原文已读取。它不会自动替代你的阅读笔记。')).to_be_visible()
        page.get_by_role('button',name='继续整理',exact=True).click();page.get_by_role('textbox',name='笔记正文',exact=True).wait_for()
    start([OUT/'note.docx'])
    editor=page.get_by_role('textbox',name='笔记正文',exact=True)
    assert editor.locator('h2').all_text_contents()==['研究问题与已有工作不足','框架图与工作流','实验分析','个人判断与论文关联'],editor.inner_html()
    assert editor.locator('img').count()==1 and editor.locator('table').count()==2
    assert editor.locator('td[colspan="2"],th[colspan="2"]').count()==1
    assert '0.31' in editor.inner_text()
    page.get_by_role('button',name='预览',exact=True).click();preview=page.locator('.source-panel .note-body');assert preview.locator('table').count()==2
    page.locator('.body-editor-heading').scroll_into_view_if_needed();page.screenshot(path=str(OUT/'desktop-editor.png'))
    editor.click();editor.press('Control+End');editor.press('Enter');editor.type('追加的手写备注 SENTINEL-EDIT。')
    page.get_by_role('button',name='插入图片',exact=True).click();page.get_by_test_id('editor-image-file').set_input_files(str(OUT/'second.png'));page.get_by_label('图注',exact=True).fill('图2 补充图片与原文页码');page.get_by_role('button',name='插入正文',exact=True).click()
    expect(editor.locator('img')).to_have_count(2)
    page.get_by_role('button',name='查看发布预览',exact=True).click();page.get_by_role('button',name='确认保存到本机',exact=True).click()
    reader=page.locator('.reader-paper .note-body');expect(reader.locator('img')).to_have_count(2);assert reader.locator('table').count()==2;assert 'SENTINEL-EDIT' in reader.inner_text()
    reader.locator('img').first.click();expect(page.get_by_role('dialog',name='查看原图')).to_be_visible();page.get_by_role('button',name='关闭',exact=True).click()
    page.reload();page.wait_for_load_state('networkidle');expect(page.locator('.reader-paper .note-body img')).to_have_count(2);assert '0.31' in page.locator('.reader-paper').inner_text()
    with page.expect_download() as download:page.get_by_role('button',name='导出笔记包',exact=True).click()
    download.value.save_as(str(OUT/'exported.zip'))
    with zipfile.ZipFile(OUT/'exported.zip') as z:
        assert len(z.namelist())==3
        exported=z.read('note.md').decode();assert 'SENTINEL-EDIT' in exported and 'colspan="2"' in exported
        assert '../assets/' not in exported and 'data:image' not in exported
    results.append('Word custom headings/order, image, ordinary and merged tables, visible edits, image insertion, lightbox, save/reload and ZIP export passed')
    start([OUT/'exported.zip']);editor=page.get_by_role('textbox',name='笔记正文',exact=True);assert editor.locator('img').count()==2 and editor.locator('table').count()==2
    assert 'SENTINEL-EDIT' in editor.inner_text();results.append('exported ZIP can be reimported without losing images or merged table')
    start([OUT/'note.md',OUT/'figure.png'],OUT/'paper.pdf');editor=page.get_by_role('textbox',name='笔记正文',exact=True)
    assert editor.locator('img').count()==1 and editor.locator('table').count()==1
    assert editor.locator('h2').all_text_contents()==['自定义第一节','自定义第二节']
    assert not page.evaluate('Boolean(window.__BAD__)')
    page.get_by_role('button',name='PDF 选页插图',exact=True).click();page.get_by_role('spinbutton',name='PDF 页码').fill('2');page.get_by_role('button',name='预览页面',exact=True).click();page.get_by_role('button',name='插入本页图片',exact=True).wait_for(state='visible');expect(page.get_by_role('button',name='插入本页图片',exact=True)).to_be_enabled(timeout=30000);page.get_by_role('button',name='插入本页图片',exact=True).click();expect(editor.locator('img')).to_have_count(2)
    assert '原文第 2 页' in editor.inner_text();results.append('Markdown with companion image, table and HTML sanitization; PDF page 2 preview and image insertion passed')
    page.wait_for_timeout(900);page.reload();page.get_by_role('button',name='恢复继续编辑',exact=True).click();editor=page.get_by_role('textbox',name='笔记正文',exact=True);expect(editor.locator('img')).to_have_count(2)
    page.locator('summary').filter(has_text='笔记图片').click();page.get_by_role('button',name='删除图片 pdf-page-2.webp',exact=True).click();expect(editor.locator('img')).to_have_count(1)
    results.append('draft retains the complete body and images; deleting an image removes its body reference')
    page.set_viewport_size({'width':390,'height':844});page.locator('.body-editor-heading').scroll_into_view_if_needed();assert page.evaluate('document.documentElement.scrollWidth <= innerWidth');page.screenshot(path=str(OUT/'mobile-editor.png'))
    results.append('390px mobile editor has no horizontal document overflow')
    start([OUT/'note.md']);page.get_by_role('button',name='查看发布预览',exact=True).click();page.get_by_role('button',name='确认保存到本机',exact=True).click();expect(page.get_by_role('alert')).to_contain_text('图片尚未绑定')
    results.append('unbound Markdown images block publication with an actionable error')
    scanned=OUT/'scanned-paper.pdf';scan=canvas.Canvas(str(scanned));scan.rect(40,40,200,200,fill=1);scan.save()
    page.goto(BASE+'#import');page.reload();page.wait_for_load_state('networkidle')
    page.get_by_test_id('note-file').set_input_files(str(OUT/'note.zip'));page.get_by_text('已解析，可继续整理',exact=True).wait_for()
    page.get_by_test_id('paper-file').set_input_files(str(scanned))
    expect(page.get_by_text('论文原文已读取。它不会自动替代你的阅读笔记。',exact=True)).to_be_visible()
    results.append('scanned original paper can accompany independent written notes; OCR is required only for note text extraction')
    bomb=OUT/'oversized-expanded.docx'
    with zipfile.ZipFile(bomb,'w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml',b' '*(31*1024*1024))
    page.goto(BASE+'#import');page.reload();page.wait_for_load_state('networkidle')
    page.get_by_test_id('note-file').set_input_files(str(bomb))
    expect(page.get_by_role('alert')).to_contain_text('展开后超过 30 MB')
    results.append('oversized compressed Word expansion rejected before document parsing')
    assert not errors,errors
    browser.close()
(OUT/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(results,ensure_ascii=False))
