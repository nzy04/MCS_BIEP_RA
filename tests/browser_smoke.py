import sys, json, pathlib, os
sys.path.insert(0, str(pathlib.Path('test-results/python-deps').resolve()))
from playwright.sync_api import sync_playwright
from docx import Document
from docx.shared import Inches
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

out=pathlib.Path('test-results').resolve();out.mkdir(exist_ok=True)
import base64
(out/'pixel.png').write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jY1kAAAAASUVORK5CYII='))
doc=Document();doc.add_heading('导入验收笔记',0);doc.add_heading('一句话总结',1);doc.add_paragraph('这是 Word 文件真实导入的内容。');doc.add_heading('核心方法',1);doc.add_paragraph('以结构化章节保留研究理解。');doc.add_picture(str(out/'pixel.png'),width=Inches(.4));doc.add_heading('我的思考',1);doc.add_paragraph('需要保留我的备注，不自动生成其他结论。');doc.save(out/'sample.docx')
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
pdf=canvas.Canvas(str(out/'sample.pdf'));pdf.drawString(60,770,'PDF import verification');pdf.drawString(60,740,'This is selectable PDF text used to validate the real extraction pipeline.');pdf.drawString(60,710,'The importer must preserve this sentence.');pdf.setFont('STSong-Light',12);pdf.drawString(60,680,'中文笔记也应当正确提取');pdf.save()
scan=canvas.Canvas(str(out/'scan.pdf'));scan.rect(40,40,200,200,fill=1);scan.save()
(out/'sample.md').write_text('# Markdown 验收\n\n## 研究背景\n保留背景内容。\n\n## 我的思考\n测试输入 <script>alert(1)</script> 不应执行。',encoding='utf-8')
results=[]
base_url=os.environ.get('TEST_BASE_URL','http://127.0.0.1:5188/')
def demo_config(route):
    route.fulfill(json={'title':'MCS_BIEP','configured':False,'authConfigured':False,'repository':'','branch':'main'})
with sync_playwright() as p:
    browser=p.chromium.launch(channel=os.environ.get('TEST_BROWSER_CHANNEL', 'msedge') or None,headless=True)
    ctx=browser.new_context(viewport={'width':1440,'height':1100},device_scale_factor=1)
    ctx.route('**/api/config',demo_config)
    def demo_data(route):
        resource=route.request.url.split('/data/',1)[1].split('?',1)[0]
        fixture=pathlib.Path('tests/fixtures/demo')/resource
        route.fulfill(body=fixture.read_text(encoding='utf-8'),content_type='application/json')
    ctx.route('**/data/**',demo_data)
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    page.goto(base_url);page.wait_for_load_state('networkidle');page.locator('.paper-card').first.wait_for()
    assert page.locator('.paper-card').count()==3
    page.screenshot(path=str(out/'desktop-library.png'),full_page=True)
    page.get_by_role('textbox',name='搜索论文').fill('SOTOPIA');assert page.locator('.paper-card').count()==1
    page.get_by_role('textbox',name='搜索论文').fill('');results.append('library: 3 demo notes; title search works')
    page.get_by_role('button',name='GitHub 登录',exact=True).click();page.get_by_role('dialog').wait_for();assert page.get_by_role('dialog').inner_text().find('本地体验版')>=0;page.get_by_role('button',name='先体验本地功能').click();results.append('unconfigured login does not simulate authorization')
    page.get_by_role('button',name='上传笔记',exact=True).click();page.get_by_test_id('note-file').set_input_files(str(out/'sample.docx'));page.get_by_text('已解析，可继续整理').wait_for(timeout=30000);page.get_by_role('button',name='继续整理',exact=True).click()
    assert page.get_by_label('笔记标题').input_value()=='导入验收笔记'
    assert '真实导入' in page.get_by_role('textbox',name='笔记正文',exact=True).inner_text()
    assert '结构化章节' in page.get_by_role('textbox',name='笔记正文',exact=True).inner_text()
    assert '我的备注' in page.get_by_role('textbox',name='笔记正文',exact=True).inner_text()
    page.locator('.metadata-panel > summary').click();page.get_by_label('分享人',exact=True).fill('验收成员');page.get_by_role('button',name='新建研究方向',exact=True).click();page.get_by_label('新研究方向',exact=True).fill('导入验收');page.get_by_role('button',name='创建方向',exact=True).click();page.get_by_role('checkbox',name='世界模型',exact=True).check()
    page.get_by_role('button',name='预览',exact=True).click();assert page.locator('.source-panel .note-body img').count()==1
    assert page.locator('.source-panel .note-body img').first.evaluate('(img)=>img.complete && img.naturalWidth>0')
    page.screenshot(path=str(out/'desktop-editor.png'),full_page=False)
    page.get_by_role('button',name='查看发布预览',exact=True).click();page.get_by_role('button',name='确认保存到本机',exact=True).click();page.get_by_text('这篇笔记仅保存在此浏览器').wait_for();page.reload();page.wait_for_load_state('networkidle');assert '导入验收笔记' in page.locator('main').inner_text();results.append('DOCX: headings, text, image, edit, save and reload persistence passed')
    stored_assets=page.evaluate("""async () => {
      const db=await new Promise((resolve,reject)=>{const r=indexedDB.open('paper-commons-v1');r.onsuccess=()=>resolve(r.result);r.onerror=()=>reject(r.error)});
      const notes=await new Promise((resolve,reject)=>{const r=db.transaction('records').objectStore('records').get('local-notes');r.onsuccess=()=>resolve(r.result);r.onerror=()=>reject(r.error)});
      db.close();return notes[0].assets.map(a=>({size:a.size,dataUrl:a.dataUrl}));
    }""")
    assert all(a['size']==len(base64.b64decode(a['dataUrl'].split(',')[1])) for a in stored_assets)
    results.append('Word image metadata matches exact upload bytes, including base64 padding')
    page.get_by_role('button',name='返回论文库',exact=True).click();assert page.locator('.paper-card').count()==4
    page.get_by_role('button',name='上传笔记',exact=True).click();page.get_by_test_id('note-file').set_input_files(str(out/'sample.pdf'));page.get_by_text('已解析，可继续整理').wait_for(timeout=45000);page.get_by_role('button',name='继续整理',exact=True).click();assert 'importer must preserve' in page.get_by_role('textbox',name='笔记正文',exact=True).inner_text();assert '中文笔记也应当正确提取' in page.get_by_role('textbox',name='笔记正文',exact=True).inner_text();results.append('PDF: actual English and Chinese text extraction passed')
    page.get_by_role('button',name='返回导入',exact=True).click();page.get_by_test_id('note-file').set_input_files(str(out/'scan.pdf'));page.get_by_role('alert').wait_for();assert '扫描件' in page.get_by_role('alert').inner_text();results.append('scan-only PDF rejected with clear OCR limitation')
    mobile_context=browser.new_context(viewport={'width':390,'height':844});mobile_context.route('**/api/config',demo_config);mobile_context.route('**/data/**',demo_data);mobile=mobile_context.new_page();mobile.goto(base_url);mobile.wait_for_load_state('networkidle');mobile.locator('.paper-card').first.wait_for();assert mobile.evaluate('document.documentElement.scrollWidth <= innerWidth');mobile.screenshot(path=str(out/'mobile-library.png'),full_page=True);results.append('mobile: 390px no horizontal overflow')
    auth_context=browser.new_context();expired={'value':False}
    def auth_fixture(route):
        path=route.request.url.split('/api/',1)[1]
        if path=='config': route.fulfill(json={'title':'MCS_BIEP','configured':True,'authConfigured':True,'repository':'test-group/test-repo','branch':'main'})
        elif path=='session': route.fulfill(json={'login':'test-member','name':'测试成员','avatar':'','canWrite':True})
        elif path=='library': route.fulfill(json={'notes':[]})
        elif path=='directions': route.fulfill(json={'directions':[]})
        elif path=='publish':
            expired['value']=True
            route.fulfill(status=401,json={'error':'请先登录 GitHub，再发布笔记。'})
        else: route.fulfill(status=404,json={'error':'Test endpoint not configured'})
    auth_context.route('**/api/**',auth_fixture)
    auth_page=auth_context.new_page();auth_page.goto(base_url);auth_page.wait_for_load_state('networkidle');auth_page.get_by_role('button',name='上传笔记',exact=True).click();auth_page.get_by_role('button',name='在线写笔记',exact=True).click();auth_page.get_by_label('笔记标题').fill('会话失效后保留的草稿');auth_page.get_by_role('textbox',name='笔记正文',exact=True).fill('服务端拒绝登录后，应显示重新登录入口并保留草稿。');auth_page.get_by_role('button',name='查看发布预览',exact=True).click();auth_page.get_by_role('button',name='确认发布到仓库',exact=True).click();auth_page.get_by_role('alert').filter(has_text='请先登录 GitHub').wait_for();assert auth_page.get_by_role('button',name='退出登录',exact=True).count()==0;assert auth_page.get_by_role('button',name='登录 GitHub',exact=True).is_visible();auth_page.wait_for_timeout(800);auth_page.reload();auth_page.get_by_role('button',name='恢复继续编辑',exact=True).wait_for();auth_page.wait_for_timeout(1000);auth_page.get_by_role('button',name='恢复继续编辑',exact=True).click();assert auth_page.get_by_label('笔记标题').input_value()=='会话失效后保留的草稿';results.append('expired session: stale identity cleared, login offered and signed-in draft recovery preserved after delay')
    assert not errors, errors
    results.append('no browser runtime errors')
    browser.close()
(out/'browser-results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=False))
